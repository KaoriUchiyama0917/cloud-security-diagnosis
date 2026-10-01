from openpyxl import load_workbook


def load_hearing_sheet(file_path):
    """
    Excel形式のヒアリングシートを読み込む。

    戻り値:
    {
        "基本情報": {
            "質問": "回答"
        },
        "権限系": {
            "MFA": {
                "質問": "回答"
            }
        }
    }

    空欄は None として扱う。
    """

    workbook = load_workbook(file_path, data_only=True)

    answers = {
        "基本情報": {},
        "権限系": {}
    }

    # ==================================================
    # 基本情報シート
    # ==================================================

    sheet = workbook["基本情報"]

    # 1行目は見出しなので2行目から読む
    for row in sheet.iter_rows(min_row=2, values_only=True):

        question = row[0]
        answer = row[1]

        if question is None:
            continue

        question = str(question).strip()

        # 空欄は None
        if answer is not None:
            answer = str(answer).strip()

            if answer == "":
                answer = None

        answers["基本情報"][question] = answer

    # ==================================================
    # 権限系シート
    # ==================================================

    sheet = workbook["権限系"]

    for row in sheet.iter_rows(min_row=2, values_only=True):

        diagnosis_item = row[0]
        question = row[1]
        answer = row[2]

        if diagnosis_item is None or question is None:
            continue

        diagnosis_item = str(diagnosis_item).strip()
        question = str(question).strip()

        # 空欄は None
        if answer is not None:
            answer = str(answer).strip()

            if answer == "":
                answer = None

        # MFAなど診断項目ごとにまとめる
        if diagnosis_item not in answers["権限系"]:
            answers["権限系"][diagnosis_item] = {}

        answers["権限系"][diagnosis_item][question] = answer

    return answers
